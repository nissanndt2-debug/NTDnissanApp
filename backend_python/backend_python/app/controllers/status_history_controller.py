from io import BytesIO
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.constants.index import EXCEL_STYLES, ROLE_IDS
from app.services.status_history_service import status_history_service
from app.utils.helpers import get_user_plant_filter


MX_TIMEZONE = ZoneInfo("America/Mexico_City")
UTC_TIMEZONE = timezone.utc


def _build_log_filters(filters: dict, user: dict | None):
    merged = dict(filters or {})

    # Carrier users can only view their provider data.
    if user and user.get("roleId") == ROLE_IDS["CARRIER"] and user.get("providerId"):
        merged["providerId"] = int(user["providerId"])

    # Keep frontend-selected plant when provided; otherwise default to user plant.
    user_plant = get_user_plant_filter(user)
    if user_plant and not merged.get("plant"):
        merged["plant"] = user_plant

    return merged


async def get_logs(filters: dict, user: dict | None = None):
    effective_filters = _build_log_filters(filters, user)
    rows = await status_history_service.search(effective_filters)

    normalized_rows: list[dict] = []
    for row in rows:
        item = dict(row)
        item["changedAt"] = _to_mx_datetime_iso(item.get("changedAt"))
        normalized_rows.append(item)

    return {"ok": True, "data": normalized_rows}


STATUS_TO_SPANISH = {
    "REPORTED": "Reportada",
    "SENT": "Nivelacion WWS",
    "DELIVERED": "Entregada a Body",
    "RECEIVED": "Recibida en Body",
    "IN_REPAIR": "En Reparacion",
    "RELEASED": "Liberada Body",
    "WTY_PENDING": "Validacion WTY",
    "WTY_RELEASED": "Liberada WTY",
    "WWS_RELEASED": "Liberada WWS",
    "ACCEPTED": "Aceptada Carrier",
    "REJECTED": "Rechazada Carrier",
    "ARCHIVED": "Archivada",
    "UNAVAILABLE": "No disponible",
}


def _to_mx_datetime(value):
    if not value:
        return ""

    parsed: datetime | None = None

    if isinstance(value, datetime):
        parsed = value

    elif isinstance(value, str):
        raw = value.strip().replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(raw)
        except ValueError:
            return value

    if parsed is None:
        return str(value)

    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC_TIMEZONE)

    return parsed.astimezone(MX_TIMEZONE).strftime("%d/%m/%Y %H:%M")


def _to_mx_datetime_iso(value):
    if not value:
        return None

    parsed: datetime | None = None

    if isinstance(value, datetime):
        parsed = value

    elif isinstance(value, str):
        raw = value.strip().replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(raw)
        except ValueError:
            return value

    if parsed is None:
        return str(value)

    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC_TIMEZONE)

    return parsed.astimezone(MX_TIMEZONE).isoformat(timespec="seconds")


async def export_logs_to_excel(filters: dict, user: dict | None = None):
    effective_filters = _build_log_filters(filters, user)
    data = await status_history_service.search(effective_filters)

    wb = Workbook()
    wb.properties.creator = EXCEL_STYLES["APP_NAME"]
    wb.properties.created = datetime.now(UTC_TIMEZONE)

    ws = wb.active
    ws.title = "Historico de Unidades"

    # Agrupar filas por unidad para generar una sola fila resumen por VIN.
    grouped: dict[int, dict] = {}
    for row in data:
        unit_id = int(row.get("unitId"))
        if unit_id not in grouped:
            grouped[unit_id] = {
                "unitId": unit_id,
                "vin": row.get("vin"),
                "market": row.get("market"),
                "lane": row.get("lane"),
                "registeredByName": row.get("registeredByName"),
                "states": {},
                "notes": [],
            }

        status = row.get("newStatus")
        if status:
            grouped[unit_id]["states"][status] = row.get("changedAt")
            if row.get("note"):
                grouped[unit_id]["notes"].append(
                    {
                        "status": status,
                        "note": row.get("note"),
                        "timestamp": row.get("changedAt"),
                    }
                )

    headers = [
        "VIN",
        "Mercado",
        "Carril",
        "Reportada (Carrier/WWS)",
        "Nivelacion (WWS)",
        "Entregada (WWS)",
        "Recibida (Body)",
        "En Reparacion (Body)",
        "Liberada Body (Body)",
        "Validacion WTY (WTY/SCM Quality)",
        "Liberada WTY (WTY/SCM Quality)",
        "Liberada WWS (WWS)",
        "Aceptada (Carrier)",
        "Rechazada (Carrier)",
        "Registrado por",
        "Nota",
    ]
    ws.append(headers)

    column_widths = [20, 14, 12, 24, 24, 24, 24, 24, 24, 30, 30, 24, 24, 24, 22, 55]
    for idx, width in enumerate(column_widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width

    thin = Side(style="thin", color=EXCEL_STYLES["BORDER_COLOR"])
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_fill = PatternFill(fill_type="solid", fgColor=EXCEL_STYLES["NISSAN_RED"])
    header_font = Font(bold=True, color=EXCEL_STYLES["WHITE"], name=EXCEL_STYLES["FONT_NAME"], size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = border

    ws.row_dimensions[1].height = 22

    ordered_units = list(grouped.values())
    for idx, unit in enumerate(ordered_units):
        states = unit["states"]

        notes_parts = []
        for item in unit["notes"]:
            label = STATUS_TO_SPANISH.get(item["status"], item["status"])
            notes_parts.append(f"[{label}] {_to_mx_datetime(item['timestamp'])}: {item['note']}")
        notes_text = "\n".join(notes_parts)

        ws.append(
            [
                unit["vin"],
                unit["market"],
                unit["lane"],
                _to_mx_datetime(states.get("REPORTED")),
                _to_mx_datetime(states.get("SENT")),
                _to_mx_datetime(states.get("DELIVERED")),
                _to_mx_datetime(states.get("RECEIVED")),
                _to_mx_datetime(states.get("IN_REPAIR")),
                _to_mx_datetime(states.get("RELEASED")),
                _to_mx_datetime(states.get("WTY_PENDING")),
                _to_mx_datetime(states.get("WTY_RELEASED")),
                _to_mx_datetime(states.get("WWS_RELEASED")),
                _to_mx_datetime(states.get("ACCEPTED")),
                _to_mx_datetime(states.get("REJECTED")),
                unit["registeredByName"],
                notes_text,
            ]
        )

        row_number = idx + 2
        row_fill = PatternFill(
            fill_type="solid",
            fgColor=EXCEL_STYLES["ROW_NORMAL"] if idx % 2 == 0 else EXCEL_STYLES["ROW_ALT"],
        )

        line_count = max(1, len(notes_parts))
        ws.row_dimensions[row_number].height = max(18, line_count * 16)

        for cell in ws[row_number]:
            cell.fill = row_fill
            cell.font = Font(name=EXCEL_STYLES["FONT_NAME"], size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
